"""The Word reader keeps hyperlink addresses (2026-09-28): the brand and category research
documents cite their facts by hyperlink, and python-docx's Paragraph.text drops the target."""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
import pytest

docx = pytest.importorskip("docx")
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

import brief_ingest  # noqa: E402


def _link(paragraph, text, url):
    """Append an external hyperlink run to `paragraph` (python-docx has no public API)."""
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def _doc(tmp_path):
    d = docx.Document()
    p = d.add_paragraph("Awareness fell to 27% (source: ")
    _link(p, "Kantar tracker", "https://example.org/kantar-2024")
    p.add_run(").")
    q = d.add_paragraph("See ")
    _link(q, "www.example.org/report/", "https://www.example.org/report/")
    m = d.add_paragraph("Contact ")
    _link(m, "a.person@example.org", "mailto:a.person@example.org")
    cell = d.add_table(rows=1, cols=2).rows[0].cells[1]
    _link(cell.paragraphs[0], "the study", "https://example.org/study")
    f = tmp_path / "research.docx"
    d.save(str(f))
    return f


def test_link_address_is_kept_after_its_text(tmp_path):
    text = brief_ingest.docx_text(_doc(tmp_path))
    assert "Awareness fell to 27% (source: Kantar tracker <https://example.org/kantar-2024>)." in text


def test_link_whose_text_is_its_address_is_not_repeated(tmp_path):
    text = brief_ingest.docx_text(_doc(tmp_path))
    assert "See www.example.org/report/\n" in text + "\n" and "<https://www.example.org/report/>" not in text
    assert "Contact a.person@example.org" in text and "mailto:" not in text


def test_links_inside_tables_keep_their_address(tmp_path):
    text = brief_ingest.docx_text(_doc(tmp_path))
    assert " | the study <https://example.org/study>" in text


def test_a_document_without_links_reads_as_before(tmp_path):
    d = docx.Document()
    d.add_paragraph("Plain paragraph.")
    d.add_table(rows=1, cols=2).rows[0].cells[0].text = "cell"
    f = tmp_path / "plain.docx"
    d.save(str(f))
    assert brief_ingest.docx_text(f) == "Plain paragraph.\ncell | "
